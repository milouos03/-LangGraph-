from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSET_DIR = ROOT / "docs" / "zero_to_agent_assets"
OUTPUT = ROOT / "docs" / "从0搭建MiniResearchAgent学习手册.docx"

FONT_BODY = "Microsoft YaHei"
FONT_CODE = "Consolas"
NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(89, 99, 110)
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
PALE_GOLD = "FFF8E8"
GREEN_FILL = "EAF5EE"


def set_run_font(run, name=FONT_BODY, size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=MUTED)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = FONT_BODY
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True

    code_style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = FONT_CODE
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CODE)
    code_style.font.size = Pt(8.3)
    code_style.font.color.rgb = RGBColor(40, 40, 40)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(8)
    code_style.paragraph_format.line_spacing = 1.0
    code_style.paragraph_format.keep_together = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("从 0 搭建 Mini Research Agent")
    set_run_font(hr, size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_real_numbering(doc: Document, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(el.get(qn("w:abstractNumId")))
        for el in numbering.findall(qn("w:abstractNum"))
        if el.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(el.get(qn("w:numId")))
        for el in numbering.findall(qn("w:num"))
        if el.get(qn("w:numId")) is not None
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if kind == "number" else "bullet")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if kind == "number" else "•")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    ppr.append(num_pr)


def add_list_item(doc, text: str, num_id: int):
    p = doc.add_paragraph()
    apply_num(p, num_id)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_para(doc, text: str, *, bold_prefix: str | None = None, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11, color=INK, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11, color=INK, italic=italic)
    return p


def add_code(doc, code: str, language="python"):
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(4)
    label.paragraph_format.space_after = Pt(2)
    rr = label.add_run(language.upper())
    set_run_font(rr, size=8.5, color=MUTED, bold=True)

    p = doc.add_paragraph(style="Code Block")
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), "2E74B5")
    borders.append(left)
    ppr.append(borders)
    r = p.add_run(code.strip("\n"))
    set_run_font(r, FONT_CODE, size=8.3, color=RGBColor(40, 40, 40))
    return p


def add_callout(doc, label: str, text: str, fill=PALE_GOLD):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Pt(8)
    p.paragraph_format.right_indent = Pt(8)
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "20")
    left.set(qn("w:color"), "2E74B5")
    borders.append(left)
    ppr.append(borders)
    r1 = p.add_run(f"{label}：")
    set_run_font(r1, size=10.5, color=NAVY, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        set_cell_margins(cell, top=100, start=120, bottom=100, end=120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=9.3, color=NAVY, bold=True)
    for row_values in rows:
        row = table.add_row()
        for idx, text in enumerate(row_values):
            cell = row.cells[idx]
            set_cell_margins(cell, top=100, start=120, bottom=100, end=120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if idx == 0 and len(headers) > 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            set_run_font(r, size=9.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _diagram_font(size: int, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def draw_arrow(draw, start, end, color="#52606D", width=5):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        sign = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - sign * 18, y2 - 10), (x2 - sign * 18, y2 + 10)]
    else:
        sign = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 10, y2 - sign * 18), (x2 + 10, y2 - sign * 18)]
    draw.polygon(points, fill=color)


def draw_box(draw, box, title, subtitle="", fill="#E8EEF5", outline="#2E74B5"):
    draw.rounded_rectangle(box, radius=24, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = box
    title_font = _diagram_font(30, True)
    sub_font = _diagram_font(22)
    tb = draw.textbbox((0, 0), title, font=title_font)
    tx = (x1 + x2 - (tb[2] - tb[0])) / 2
    ty = y1 + 28
    draw.text((tx, ty), title, fill="#0B2545", font=title_font)
    if subtitle:
        sb = draw.multiline_textbbox((0, 0), subtitle, font=sub_font, spacing=6, align="center")
        sx = (x1 + x2 - (sb[2] - sb[0])) / 2
        sy = y1 + 82
        draw.multiline_text((sx, sy), subtitle, fill="#46515D", font=sub_font, spacing=6, align="center")


def create_diagrams():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    img = Image.new("RGB", (1700, 720), "white")
    draw = ImageDraw.Draw(img)
    title_font = _diagram_font(38, True)
    draw.text((70, 45), "从 0 到可用 Demo：四阶段学习路线", fill="#0B2545", font=title_font)
    boxes = [
        (70, 200, 400, 470),
        (480, 200, 810, 470),
        (890, 200, 1220, 470),
        (1300, 200, 1630, 470),
    ]
    labels = [
        ("阶段 1", "单次模型调用\n确认 API 与配置"),
        ("阶段 2", "LangGraph 编排\n状态、节点、路由"),
        ("阶段 3", "FastAPI 后端\n把工作流变成接口"),
        ("阶段 4", "Vue 前端\n完成端到端 Demo"),
    ]
    fills = ["#EAF5EE", "#E8EEF5", "#FFF8E8", "#F4ECF7"]
    for box, (title, sub), fill in zip(boxes, labels, fills):
        draw_box(draw, box, title, sub, fill=fill)
    for idx in range(3):
        draw_arrow(draw, (boxes[idx][2] + 8, 335), (boxes[idx + 1][0] - 10, 335))
    note_font = _diagram_font(24)
    draw.text((70, 560), "原则：每一阶段都必须可以独立运行、独立验证，再进入下一阶段。", fill="#52606D", font=note_font)
    img.save(ASSET_DIR / "learning-roadmap.png", quality=95)

    img = Image.new("RGB", (1700, 900), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 40), "Mini Research Agent 最小架构", fill="#0B2545", font=title_font)
    top = [
        ((80, 180, 370, 360), "Vue", "用户输入\n展示结果"),
        ((510, 180, 800, 360), "FastAPI", "请求校验\n调用工作流"),
        ((940, 180, 1230, 360), "LangGraph", "状态流转\n条件路由"),
        ((1370, 180, 1620, 360), "LLM", "通义千问\n角色推理"),
    ]
    for box, title, sub in top:
        draw_box(draw, box, title, sub)
    for idx in range(3):
        draw_arrow(draw, (top[idx][0][2] + 8, 270), (top[idx + 1][0][0] - 10, 270))

    lower = [
        ((650, 560, 990, 760), "state.py", "共享状态"),
        ((1080, 560, 1420, 760), "nodes.py", "节点业务逻辑"),
    ]
    for box, title, sub in lower:
        draw_box(draw, box, title, sub, fill="#F4F6F9", outline="#52606D")
    draw_arrow(draw, (1085, 375), (820, 550))
    draw_arrow(draw, (1130, 375), (1250, 550))
    img.save(ASSET_DIR / "mini-architecture.png", quality=95)

    img = Image.new("RGB", (1700, 1050), "white")
    draw = ImageDraw.Draw(img)
    draw.text((70, 40), "LangGraph 工作流：简单问题与研究问题分流", fill="#0B2545", font=title_font)
    nodes = {
        "start": (680, 130, 1020, 270),
        "intent": (680, 340, 1020, 500),
        "direct": (160, 600, 500, 780),
        "plan": (640, 600, 980, 780),
        "research": (1100, 600, 1440, 780),
        "writer": (1100, 860, 1440, 1010),
        "end": (160, 860, 500, 1010),
    }
    draw_box(draw, nodes["start"], "START", "用户问题", fill="#EAF5EE")
    draw_box(draw, nodes["intent"], "intent", "判断 direct / research")
    draw_box(draw, nodes["direct"], "direct_answer", "直接回答", fill="#FFF8E8")
    draw_box(draw, nodes["plan"], "plan", "拆分研究步骤")
    draw_box(draw, nodes["research"], "research", "整理与分析")
    draw_box(draw, nodes["writer"], "writer", "生成最终报告", fill="#F4ECF7")
    draw_box(draw, nodes["end"], "END", "返回 final", fill="#EAF5EE")
    draw_arrow(draw, (850, 280), (850, 330))
    draw_arrow(draw, (675, 485), (340, 590))
    draw_arrow(draw, (850, 510), (810, 590))
    draw_arrow(draw, (990, 690), (1090, 690))
    draw_arrow(draw, (1270, 790), (1270, 850))
    draw_arrow(draw, (1090, 930), (510, 930))
    draw_arrow(draw, (330, 790), (330, 850))
    small = _diagram_font(23, True)
    draw.text((450, 515), "direct", fill="#7A5A00", font=small)
    draw.text((875, 530), "research", fill="#2E74B5", font=small)
    img.save(ASSET_DIR / "workflow-routing.png", quality=95)


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run()
    shape = r.add_picture(str(path), width=Inches(6.25))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cr = cap.add_run(caption)
    set_run_font(cr, size=9.5, color=MUTED, italic=True)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(90)
    p.paragraph_format.space_after = Pt(18)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("实战学习手册")
    set_run_font(r, size=12, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("从 0 搭建 Mini Research Agent")
    set_run_font(r, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Python + LangGraph + FastAPI + Vue 3")
    set_run_font(r, size=15, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(36)
    r = p.add_run("用最小可运行版本理解多智能体项目的真正骨架")
    set_run_font(r, size=11.5, color=MUTED, italic=True)

    add_callout(
        doc,
        "学习目标",
        "独立完成一个支持“简单问答 / 研究型回答”分流的智能体 Demo，"
        "并理解配置、状态、节点、图编排、API 和前端之间的关系。",
        fill=LIGHT_BLUE,
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(55)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("建议学习周期：7 天  ·  适合已有 Python 基础的初学者")
    set_run_font(r, size=10, color=MUTED)
    p.add_run().add_break(WD_BREAK.PAGE)


def build_document():
    create_diagrams()
    doc = Document()
    configure_document(doc)
    bullet_id = add_real_numbering(doc, "bullet")
    number_id = add_real_numbering(doc, "number")
    add_cover(doc)

    doc.add_heading("导读：先缩小问题，再开始编码", level=1)
    add_para(
        doc,
        "原项目同时包含前端、后端、LangGraph、多 Agent、Web Search、RAG、记忆、"
        "Redis、PostgreSQL、Milvus 和 SSE。它不是不适合学习，而是不适合一次性学习。"
        "本手册把目标压缩为一个可独立完成的最小版本。",
    )
    add_callout(
        doc,
        "最小闭环",
        "用户提问 → 判断问题类型 → 直接回答或制定计划 → 分析 → 生成结果。",
        fill=GREEN_FILL,
    )
    add_figure(doc, ASSET_DIR / "learning-roadmap.png", "图 1  四阶段学习路线")

    doc.add_heading("你最终会得到什么", level=2)
    for item in (
        "一个可以在命令行中运行的 LangGraph 智能体。",
        "一个提供 POST /api/research 接口的 FastAPI 后端。",
        "一个可以发送问题并展示回答的 Vue 3 页面。",
        "一套能够继续加入搜索、RAG、SSE 和记忆的清晰骨架。",
    ):
        add_list_item(doc, item, bullet_id)

    doc.add_heading("第一版主动不做的功能", level=2)
    add_para(doc, "下面这些能力不是没用，而是应该在最小闭环稳定后再加入：")
    for item in (
        "真实网络搜索和引用校验",
        "Milvus 向量数据库与文档导入",
        "Redis / PostgreSQL / 长期记忆",
        "多轮反思循环和并行检索",
        "SSE 进度流与复杂 Markdown 渲染",
    ):
        add_list_item(doc, item, bullet_id)

    doc.add_heading("里程碑与验收标准", level=2)
    add_table(
        doc,
        ["阶段", "交付结果", "验收标准"],
        [
            ["1", "单次模型调用", "终端能打印模型回答"],
            ["2", "LangGraph 工作流", "简单问题和复杂问题走不同路径"],
            ["3", "FastAPI 接口", "Swagger 中可以提交问题并得到 JSON"],
            ["4", "Vue 页面", "浏览器中完成一次端到端问答"],
        ],
        [900, 3300, 5160],
    )

    doc.add_heading("第 1 章  先建立正确的项目心智模型", level=1)
    add_figure(doc, ASSET_DIR / "mini-architecture.png", "图 2  Mini Research Agent 最小架构")
    doc.add_heading("五个核心概念", level=2)
    concepts = [
        ("配置 Config", "保存 API Key、模型名和接口地址，避免把环境差异写进业务代码。"),
        ("状态 State", "工作流中共享的数据对象，保存 query、route、plan 和 final。"),
        ("节点 Node", "一个可执行步骤。它读取 State，并返回需要更新的字段。"),
        ("图 Graph", "注册节点、连接执行顺序，并根据状态进行条件路由。"),
        ("接口 API", "把图的执行能力封装为 HTTP 服务，供 Vue 或其他客户端调用。"),
    ]
    for idx, (name, desc) in enumerate(concepts, 1):
        add_para(doc, f"{idx}. {name}：{desc}", bold_prefix=f"{idx}. {name}：")

    add_callout(
        doc,
        "一句话记忆",
        "Graph 控制顺序，Node 执行业务，State 搬运数据，API 对外提供能力，Vue 负责交互。",
    )

    doc.add_heading("第 2 章  创建项目与开发环境", level=1)
    doc.add_heading("2.1 前置条件", level=2)
    for item in (
        "Python 3.10 或更高版本",
        "Node.js 20 或更高版本",
        "一个可用的百炼 API Key",
        "VS Code 或其他 Python 编辑器",
    ):
        add_list_item(doc, item, bullet_id)

    doc.add_heading("2.2 创建目录和虚拟环境", level=2)
    add_code(
        doc,
        r"""
mkdir agent_demo
cd agent_demo

python -m venv .venv
.venv\Scripts\Activate.ps1
""",
        "powershell",
    )

    doc.add_heading("2.3 安装依赖", level=2)
    add_code(
        doc,
        """
langchain
langchain-openai
langgraph
fastapi
uvicorn[standard]
python-dotenv
pydantic
""",
        "requirements.txt",
    )
    add_code(doc, "pip install -r requirements.txt", "powershell")

    doc.add_heading("2.4 最终目录结构", level=2)
    add_code(
        doc,
        """
agent_demo/
├── .env
├── .gitignore
├── requirements.txt
├── main.py
├── app/
│   ├── __init__.py
│   ├── config.py
│   ├── state.py
│   ├── agents.py
│   ├── nodes.py
│   ├── graph.py
│   └── api.py
└── frontend/
""",
        "text",
    )

    doc.add_heading("2.5 配置密钥", level=2)
    add_code(
        doc,
        """
DASHSCOPE_API_KEY=替换为你的APIKey
DASHSCOPE_BASE_URL=https://dashscope.aliyuncs.com/compatible-mode/v1
MODEL_NAME=qwen-plus
""",
        ".env",
    )
    add_code(doc, ".env\n.venv/\n__pycache__/\nnode_modules/\n", ".gitignore")
    add_callout(doc, "安全提醒", "不要把 .env 上传到 Git，也不要把 API Key 直接写进 Python 文件。")

    doc.add_heading("第 3 章  阶段一：跑通第一次模型调用", level=1)
    doc.add_heading("3.1 编写配置模块", level=2)
    add_code(
        doc,
        """
import os

from dotenv import load_dotenv


load_dotenv()


class Settings:
    api_key = os.getenv("DASHSCOPE_API_KEY", "")
    base_url = os.getenv(
        "DASHSCOPE_BASE_URL",
        "https://dashscope.aliyuncs.com/compatible-mode/v1",
    )
    model_name = os.getenv("MODEL_NAME", "qwen-plus")


settings = Settings()
""",
        "app/config.py",
    )

    doc.add_heading("3.2 构建模型", level=2)
    add_code(
        doc,
        """
from langchain_openai import ChatOpenAI

from app.config import settings


def build_model(temperature: float = 0.2) -> ChatOpenAI:
    if not settings.api_key:
        raise RuntimeError("DASHSCOPE_API_KEY 未配置")

    return ChatOpenAI(
        model=settings.model_name,
        api_key=settings.api_key,
        base_url=settings.base_url,
        temperature=temperature,
    )
""",
        "app/agents.py",
    )

    doc.add_heading("3.3 完成第一次调用", level=2)
    add_code(
        doc,
        """
from app.agents import build_model


def main():
    model = build_model()
    response = model.invoke("请用一句话解释什么是智能体。")
    print(response.content)


if __name__ == "__main__":
    main()
""",
        "main.py",
    )
    add_code(doc, "python main.py", "powershell")
    add_callout(
        doc,
        "本章验收",
        "终端能够输出正常中文回答。若这一步失败，不要继续写 LangGraph，先排查密钥、网络和模型名。",
        fill=GREEN_FILL,
    )

    doc.add_heading("第 4 章  阶段二：用 LangGraph 编排多个角色", level=1)
    add_figure(doc, ASSET_DIR / "workflow-routing.png", "图 3  条件路由工作流")

    doc.add_heading("4.1 定义共享状态", level=2)
    add_para(
        doc,
        "节点之间不直接互相调用，而是通过 State 交换数据。每个节点读取完整状态，"
        "只返回自己负责更新的字段。",
    )
    add_code(
        doc,
        """
from typing import TypedDict


class AgentState(TypedDict):
    query: str
    route: str
    plan: str
    research: str
    final: str


def create_initial_state(query: str) -> AgentState:
    return {
        "query": query,
        "route": "",
        "plan": "",
        "research": "",
        "final": "",
    }
""",
        "app/state.py",
    )

    doc.add_heading("4.2 实现节点", level=2)
    add_para(doc, "先实现一个通用模型调用函数，再实现四个角色节点。")
    add_code(
        doc,
        """
import json
import re

from langchain_core.messages import HumanMessage, SystemMessage

from app.agents import build_model
from app.state import AgentState


model = build_model()


def invoke_agent(system_prompt: str, user_prompt: str) -> str:
    response = model.invoke(
        [
            SystemMessage(content=system_prompt),
            HumanMessage(content=user_prompt),
        ]
    )
    return str(response.content)


def extract_json(text: str) -> dict:
    cleaned = re.sub(r"^```json\\s*|\\s*```$", "", text.strip())
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        return {}
""",
        "app/nodes.py（公共函数）",
    )
    add_code(
        doc,
        """
def intent_node(state: AgentState) -> dict:
    result = invoke_agent(
        system_prompt=(
            "你是意图分类器。判断问题应该直接回答，"
            "还是需要研究。只输出 JSON："
            '{"route":"direct或research"}'
        ),
        user_prompt=state["query"],
    )

    data = extract_json(result)
    route = data.get("route", "research")
    if route not in {"direct", "research"}:
        route = "research"
    return {"route": route}


def direct_answer_node(state: AgentState) -> dict:
    answer = invoke_agent(
        system_prompt="你是简洁、准确的问答助手。",
        user_prompt=state["query"],
    )
    return {"final": answer}
""",
        "app/nodes.py（意图与直答）",
    )
    add_code(
        doc,
        """
def plan_node(state: AgentState) -> dict:
    plan = invoke_agent(
        system_prompt=(
            "你是研究规划师。请把用户问题拆成 3 个以内的"
            "研究步骤，使用简洁的编号列表。"
        ),
        user_prompt=state["query"],
    )
    return {"plan": plan}


def research_node(state: AgentState) -> dict:
    research = invoke_agent(
        system_prompt=(
            "你是研究分析师。当前没有联网工具，请根据已有"
            "知识分析问题，并明确说明不确定信息。"
        ),
        user_prompt=(
            f"用户问题：\\n{state['query']}\\n\\n"
            f"研究计划：\\n{state['plan']}"
        ),
    )
    return {"research": research}


def writer_node(state: AgentState) -> dict:
    final = invoke_agent(
        system_prompt=(
            "你是报告撰写者。根据问题、计划和分析结果，"
            "生成结构清楚、不过度冗长的最终回答。"
        ),
        user_prompt=(
            f"用户问题：\\n{state['query']}\\n\\n"
            f"研究计划：\\n{state['plan']}\\n\\n"
            f"分析材料：\\n{state['research']}"
        ),
    )
    return {"final": final}
""",
        "app/nodes.py（规划、研究与写作）",
    )
    add_callout(
        doc,
        "关键理解",
        "多 Agent 不一定意味着多个不同模型。第一版可以共享同一个模型对象，"
        "通过不同的 System Prompt 让节点扮演不同角色。",
    )

    doc.add_heading("4.3 编排图与条件路由", level=2)
    add_code(
        doc,
        """
from langgraph.graph import END, START, StateGraph

from app.nodes import (
    direct_answer_node,
    intent_node,
    plan_node,
    research_node,
    writer_node,
)
from app.state import AgentState


def route_after_intent(state: AgentState) -> str:
    if state["route"] == "direct":
        return "direct"
    return "research"


def build_graph():
    workflow = StateGraph(AgentState)

    workflow.add_node("intent", intent_node)
    workflow.add_node("direct_answer", direct_answer_node)
    workflow.add_node("plan", plan_node)
    workflow.add_node("research", research_node)
    workflow.add_node("writer", writer_node)

    workflow.add_edge(START, "intent")
    workflow.add_conditional_edges(
        "intent",
        route_after_intent,
        {
            "direct": "direct_answer",
            "research": "plan",
        },
    )
    workflow.add_edge("direct_answer", END)
    workflow.add_edge("plan", "research")
    workflow.add_edge("research", "writer")
    workflow.add_edge("writer", END)

    return workflow.compile()


agent_graph = build_graph()
""",
        "app/graph.py",
    )

    doc.add_heading("4.4 在命令行运行工作流", level=2)
    add_code(
        doc,
        """
from app.graph import agent_graph
from app.state import create_initial_state


def main():
    while True:
        query = input("\\n你：").strip()
        if query.lower() in {"exit", "quit"}:
            break
        if not query:
            continue

        initial_state = create_initial_state(query)
        result = agent_graph.invoke(initial_state)
        print("\\nAI：")
        print(result["final"])


if __name__ == "__main__":
    main()
""",
        "main.py",
    )

    doc.add_heading("4.5 测试两条路径", level=2)
    add_table(
        doc,
        ["测试问题", "预期路径"],
        [
            ["你好，你是谁？", "intent → direct_answer → END"],
            ["请分析企业应该选择单智能体还是多智能体方案", "intent → plan → research → writer → END"],
        ],
        [5000, 4360],
    )
    add_callout(doc, "本章验收", "两个问题能够走不同路径，并且都能得到 final。", fill=GREEN_FILL)

    doc.add_heading("第 5 章  观察状态如何流动", level=1)
    add_para(
        doc,
        "调试 LangGraph 时，不要只看最终答案。先观察每个节点向 State 写入了什么。",
    )
    add_code(
        doc,
        """
initial_state = create_initial_state(query)

for event in agent_graph.stream(
    initial_state,
    stream_mode="updates",
):
    print("节点更新：", event)
""",
        "python",
    )
    add_para(doc, "典型输出：")
    add_code(
        doc,
        """
{'intent': {'route': 'research'}}
{'plan': {'plan': '1. ...'}}
{'research': {'research': '...'}}
{'writer': {'final': '...'}}
""",
        "text",
    )
    doc.add_heading("调试时依次检查", level=2)
    for item in (
        "当前节点收到的 state 是否包含上一节点的结果。",
        "节点是否只返回需要更新的字段。",
        "route_after_intent() 的返回值是否和映射表中的键一致。",
        "模型是否真的只输出 JSON；解析失败时是否有兜底值。",
        "最终路径是否一定可以到达 END。",
    ):
        add_list_item(doc, item, number_id)

    doc.add_heading("第 6 章  阶段三：把工作流封装成 FastAPI", level=1)
    add_code(
        doc,
        """
from fastapi import FastAPI
from pydantic import BaseModel, Field

from app.graph import agent_graph
from app.state import create_initial_state


app = FastAPI(title="Mini Agent Demo")


class ResearchRequest(BaseModel):
    query: str = Field(min_length=1, max_length=5000)


class ResearchResponse(BaseModel):
    query: str
    route: str
    final: str


@app.get("/health")
async def health():
    return {"status": "ok"}


@app.post("/api/research", response_model=ResearchResponse)
async def research(payload: ResearchRequest):
    initial_state = create_initial_state(payload.query)
    result = agent_graph.invoke(initial_state)

    return ResearchResponse(
        query=payload.query,
        route=result["route"],
        final=result["final"],
    )
""",
        "app/api.py",
    )
    add_code(doc, "uvicorn app.api:app --reload --port 8000", "powershell")
    add_para(doc, "打开 http://127.0.0.1:8000/docs，在 Swagger 中测试：")
    add_code(
        doc,
        """
{
  "query": "请分析多智能体系统的优缺点"
}
""",
        "json",
    )
    add_callout(
        doc,
        "同步接口说明",
        "这一版会等待整个工作流完成后再返回。先确保同步接口稳定，再升级成 SSE。",
    )

    doc.add_heading("第 7 章  阶段四：接入 Vue 3 前端", level=1)
    doc.add_heading("7.1 创建项目", level=2)
    add_code(
        doc,
        """
npm create vite@latest frontend -- --template vue-ts
cd frontend
npm install
npm run dev
""",
        "powershell",
    )

    doc.add_heading("7.2 配置开发代理", level=2)
    add_code(
        doc,
        """
import { defineConfig } from 'vite'
import vue from '@vitejs/plugin-vue'

export default defineConfig({
  plugins: [vue()],
  server: {
    port: 5173,
    proxy: {
      '/api': {
        target: 'http://127.0.0.1:8000',
        changeOrigin: true,
      },
    },
  },
})
""",
        "frontend/vite.config.ts",
    )

    doc.add_heading("7.3 编写最小页面", level=2)
    add_code(
        doc,
        """
<script setup lang="ts">
import { ref } from 'vue'

const query = ref('')
const answer = ref('')
const route = ref('')
const loading = ref(false)
const error = ref('')

async function submit() {
  const text = query.value.trim()
  if (!text || loading.value) return

  loading.value = true
  answer.value = ''
  route.value = ''
  error.value = ''

  try {
    const response = await fetch('/api/research', {
      method: 'POST',
      headers: { 'Content-Type': 'application/json' },
      body: JSON.stringify({ query: text }),
    })

    if (!response.ok) {
      throw new Error(`请求失败：${response.status}`)
    }

    const data = await response.json()
    route.value = data.route
    answer.value = data.final
  } catch (err) {
    error.value = err instanceof Error ? err.message : '未知错误'
  } finally {
    loading.value = false
  }
}
</script>

<template>
  <main>
    <h1>Mini Research Agent</h1>
    <textarea v-model="query" rows="6" />
    <button :disabled="loading" @click="submit">
      {{ loading ? '处理中...' : '发送' }}
    </button>
    <p v-if="route">执行路径：{{ route }}</p>
    <p v-if="error">{{ error }}</p>
    <article v-if="answer">
      <h2>回答</h2>
      <pre>{{ answer }}</pre>
    </article>
  </main>
</template>
""",
        "frontend/src/App.vue",
    )
    add_callout(
        doc,
        "端到端验收",
        "同时运行 FastAPI 和 Vite，在浏览器输入问题，页面显示 route 与 final。",
        fill=GREEN_FILL,
    )

    doc.add_heading("第 8 章  常见问题与排查顺序", level=1)
    add_table(
        doc,
        ["现象", "优先检查", "处理方式"],
        [
            ["401 / 403", "API Key", "检查 .env、额度和密钥权限"],
            ["模型找不到", "MODEL_NAME", "改成账号已开通的模型名"],
            ["ModuleNotFoundError", "虚拟环境", "激活 .venv 后重新安装 requirements"],
            ["总是走 research", "意图 JSON", "打印模型原始输出，检查解析与兜底"],
            ["图无法编译", "节点和边名称", "检查 add_node 与映射表是否一致"],
            ["前端 404", "Vite proxy", "确认 /api 代理目标端口为 8000"],
            ["前端一直等待", "后端日志", "先用 Swagger 确认接口是否可用"],
        ],
        [2100, 2500, 4760],
    )
    doc.add_heading("推荐排查顺序", level=2)
    for item in (
        "单独调用模型，确认密钥和网络。",
        "单独调用每个节点，确认输入输出。",
        "运行 agent_graph.invoke()，确认图可达 END。",
        "用 Swagger 调用 FastAPI。",
        "最后再排查 Vue 和代理。",
    ):
        add_list_item(doc, item, number_id)

    doc.add_heading("第 9 章  第二阶段如何逐步增强", level=1)
    add_para(doc, "最小 Demo 完成后，建议严格按照下面的顺序扩展。")
    add_table(
        doc,
        ["顺序", "能力", "新增状态字段", "解决的问题"],
        [
            ["1", "真实 Web Search", "sources", "回答缺少实时证据"],
            ["2", "证据分析节点", "analysis / citations", "模型可能编造来源"],
            ["3", "SSE 流式进度", "无必需新增", "长任务缺少过程反馈"],
            ["4", "RAG", "local_sources", "无法使用内部资料"],
            ["5", "短期记忆", "thread_id", "多轮对话不能延续"],
            ["6", "长期记忆", "user_id / profile", "跨会话偏好无法保存"],
        ],
        [700, 2100, 2600, 3960],
    )

    doc.add_heading("9.1 加入真实搜索", level=2)
    add_para(
        doc,
        "不要让模型假装搜索。Python 代码负责调用真实 API，把结果写进 sources；"
        "分析节点只能根据 sources 得出结论。",
    )
    add_code(
        doc,
        """
class AgentState(TypedDict):
    query: str
    route: str
    plan: str
    sources: list[dict]
    analysis: str
    final: str
""",
        "python",
    )

    doc.add_heading("9.2 加入 RAG", level=2)
    add_para(
        doc,
        "RAG 保存项目文档和知识库；记忆保存对话和用户偏好。两者用途不同。"
        "第一版可以先用轻量本地向量库验证切块、Embedding 和召回逻辑，稳定后再换 Milvus。",
    )

    doc.add_heading("9.3 加入 SSE", level=2)
    add_para(doc, "普通接口稳定后，再让后端依次发送阶段事件：")
    add_code(
        doc,
        """
{"type": "phase", "node": "plan", "message": "正在制定计划"}
{"type": "phase", "node": "research", "message": "正在整理资料"}
{"type": "final", "final": "最终报告"}
""",
        "jsonl",
    )

    doc.add_heading("RAG、记忆和 Checkpointer 的区别", level=2)
    add_table(
        doc,
        ["能力", "保存内容", "主要用途"],
        [
            ["RAG", "文档与知识片段", "让 Agent 使用私有知识"],
            ["短期记忆", "当前会话消息", "保持多轮上下文"],
            ["长期记忆", "用户偏好和历史任务", "跨会话延续"],
            ["Checkpointer", "工作流执行状态", "恢复或持续执行 LangGraph"],
        ],
        [1800, 3300, 4260],
    )

    doc.add_heading("第 10 章  7 天实战计划", level=1)
    add_table(
        doc,
        ["天数", "学习任务", "当天交付物", "自测问题"],
        [
            ["Day 1", "环境、配置、模型调用", "模型调用脚本", "为什么配置不应写死？"],
            ["Day 2", "State 与 Node", "state.py、nodes.py", "节点为什么返回局部字段？"],
            ["Day 3", "Graph 与条件路由", "graph.py", "条件边的返回值如何匹配？"],
            ["Day 4", "调试与异常兜底", "可观察的 CLI", "JSON 解析失败怎么办？"],
            ["Day 5", "FastAPI", "可用的 /api/research", "Schema 解决什么问题？"],
            ["Day 6", "Vue 与代理", "端到端网页", "为什么开发时使用 proxy？"],
            ["Day 7", "重构、日志、复盘", "README 与架构图", "每个文件能否一句话解释？"],
        ],
        [900, 2600, 2700, 3160],
    )

    doc.add_heading("毕业检查清单", level=2)
    checklist = [
        ["□", "可以解释 Config、State、Node、Graph 和 API 的关系"],
        ["□", "可以独立添加一个新节点并把它连接到图中"],
        ["□", "可以打印每个节点的状态更新"],
        ["□", "可以判断一个错误发生在模型、工作流、后端还是前端"],
        ["□", "可以从浏览器完成一次端到端问答"],
        ["□", "知道何时应该加入搜索、RAG、SSE 和记忆"],
    ]
    add_table(doc, ["完成", "验收项"], checklist, [1000, 8360])

    doc.add_heading("附录 A  常用启动命令", level=1)
    add_code(
        doc,
        r"""
# 1. 激活后端虚拟环境
.venv\Scripts\Activate.ps1

# 2. 运行命令行版本
python main.py

# 3. 运行 FastAPI
uvicorn app.api:app --reload --port 8000

# 4. 运行 Vue（另开一个终端）
cd frontend
npm run dev
""",
        "powershell",
    )

    doc.add_heading("附录 B  核心术语速查", level=1)
    add_table(
        doc,
        ["术语", "含义"],
        [
            ["LLM", "负责理解和生成文本的大语言模型"],
            ["Agent", "由模型、提示词、工具和运行规则组成的执行角色"],
            ["State", "工作流节点共享的数据结构"],
            ["Node", "读取状态并返回状态更新的函数"],
            ["Edge", "节点之间的固定执行关系"],
            ["Conditional Edge", "根据状态动态选择下一节点"],
            ["Prompt", "给模型的角色说明、任务和输出约束"],
            ["RAG", "先检索相关资料，再让模型基于资料回答"],
            ["SSE", "服务器持续向浏览器发送单向事件流"],
            ["Checkpointer", "保存 LangGraph 的中间执行状态"],
        ],
        [2200, 7160],
    )

    doc.add_heading("附录 C  三个进阶练习", level=1)
    exercises = [
        ("练习 1：增加翻译路由", "让 intent 支持 translate，并新增 translate_node。"),
        ("练习 2：增加简单计算工具", "识别数学问题后调用 Python 函数，而不是让模型心算。"),
        ("练习 3：增加研究进度", "用 graph.stream() 在终端展示当前执行节点。"),
    ]
    for title, desc in exercises:
        doc.add_heading(title, level=2)
        add_para(doc, desc)
        add_para(doc, "验收要求：新增能力不影响 direct 和 research 两条已有路径。", italic=True)

    doc.add_heading("结语", level=1)
    add_para(
        doc,
        "你真正要掌握的不是某一个框架函数，而是把复杂系统拆成可验证的小闭环。"
        "每加入一个组件，都应该能明确回答：它解决了上一版的什么问题？"
        "如果上一版还没有遇到这个问题，就暂时不要加入这个组件。",
    )
    add_callout(
        doc,
        "最终记忆",
        "先让一条最短链路可靠运行，再让它变强；先让状态看得见，再让流程变复杂。",
        fill=GREEN_FILL,
    )

    doc.core_properties.title = "从0搭建Mini Research Agent学习手册"
    doc.core_properties.subject = "Python、LangGraph、FastAPI 与 Vue 3 实战入门"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Agent, LangGraph, FastAPI, Vue, Python"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
